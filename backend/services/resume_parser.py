from __future__ import annotations

import io
import re

# Common tech and professional skills to look for in resumes
KNOWN_SKILLS: list[str] = [
    # Languages
    "Python", "JavaScript", "TypeScript", "Java", "C++", "C#", "Go", "Rust",
    "Ruby", "PHP", "Swift", "Kotlin", "Scala", "R", "MATLAB", "Bash", "Shell",
    "Perl", "Haskell", "Lua", "Elixir", "Clojure",
    # Web frameworks / backends
    "FastAPI", "Django", "Flask", "Spring", "Spring Boot", "Express", "NestJS",
    "Rails", "Laravel", "ASP.NET", "Gin", "Echo", "Fiber",
    # Frontend
    "React", "Vue", "Angular", "Svelte", "Next.js", "Nuxt", "Redux",
    "Tailwind", "Bootstrap", "HTML", "CSS", "SASS", "LESS", "Webpack", "Vite",
    # Databases
    "PostgreSQL", "MySQL", "SQLite", "MongoDB", "Redis", "Elasticsearch",
    "Cassandra", "DynamoDB", "Snowflake", "BigQuery", "Oracle", "MSSQL",
    "Neo4j", "InfluxDB", "CouchDB",
    # Cloud / DevOps
    "AWS", "Azure", "GCP", "Google Cloud", "Docker", "Kubernetes", "Terraform",
    "Ansible", "Helm", "CI/CD", "Jenkins", "GitHub Actions", "GitLab CI",
    "CircleCI", "ArgoCD", "Prometheus", "Grafana", "Datadog", "New Relic",
    # AI / ML
    "TensorFlow", "PyTorch", "scikit-learn", "Keras", "Hugging Face",
    "LangChain", "OpenAI", "Anthropic", "RAG", "LLM", "NLP", "Computer Vision",
    "Pandas", "NumPy", "SciPy", "Matplotlib", "Seaborn", "Spark", "MLflow",
    # APIs / Protocols
    "REST", "GraphQL", "gRPC", "WebSockets", "OpenAPI", "Swagger",
    "OAuth", "JWT", "SAML", "LDAP",
    # Tools / Practices
    "Git", "Linux", "Agile", "Scrum", "Kanban", "TDD", "BDD", "Microservices",
    "Kafka", "RabbitMQ", "Celery", "Airflow", "dbt",
    # Soft / domain
    "Machine Learning", "Data Science", "Data Engineering", "DevOps",
    "Site Reliability", "SRE", "Product Management", "System Design",
]


def parse_resume(file_bytes: bytes, filename: str) -> str:
    """
    Parse resume text from uploaded file bytes.

    Supports PDF (via pdfplumber), DOCX (via python-docx), and plain TXT.
    Returns the extracted text as a single string.
    """
    lower_name = filename.lower()

    if lower_name.endswith(".pdf"):
        return _parse_pdf(file_bytes)
    elif lower_name.endswith(".docx"):
        return _parse_docx(file_bytes)
    elif lower_name.endswith(".txt"):
        return _parse_txt(file_bytes)
    else:
        # Attempt plain text as a last resort
        try:
            return file_bytes.decode("utf-8", errors="replace")
        except Exception as exc:
            raise ValueError(
                f"Unsupported file type for '{filename}'. "
                "Please upload a PDF, DOCX, or TXT file."
            ) from exc


def _parse_pdf(file_bytes: bytes) -> str:
    """Extract text from a PDF using pdfplumber."""
    import pdfplumber  # type: ignore

    text_parts: list[str] = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)

    full_text = "\n".join(text_parts).strip()
    if not full_text:
        raise ValueError(
            "No text could be extracted from the PDF. "
            "The file may be image-based or empty."
        )
    return full_text


def _parse_docx(file_bytes: bytes) -> str:
    """Extract text from a DOCX file using python-docx."""
    from docx import Document  # type: ignore

    doc = Document(io.BytesIO(file_bytes))
    paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]

    # Also capture text from tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    paragraphs.append(cell_text)

    full_text = "\n".join(paragraphs).strip()
    if not full_text:
        raise ValueError(
            "No text could be extracted from the DOCX file. The document may be empty."
        )
    return full_text


def _parse_txt(file_bytes: bytes) -> str:
    """Decode a plain-text file."""
    # Try UTF-8 first, fall back to latin-1
    for encoding in ("utf-8", "latin-1", "cp1252"):
        try:
            text = file_bytes.decode(encoding)
            if text.strip():
                return text.strip()
        except UnicodeDecodeError:
            continue

    raise ValueError("Could not decode the TXT file. Please ensure it uses UTF-8 encoding.")


def extract_skills_preview(text: str, limit: int = 10) -> list[str]:
    """
    Return up to `limit` recognised skill keywords found in `text`.

    Uses case-insensitive whole-word matching against KNOWN_SKILLS.
    """
    found: list[str] = []
    text_lower = text.lower()

    for skill in KNOWN_SKILLS:
        # Build a pattern that matches the skill as a whole word / token
        pattern = r"(?<![a-zA-Z0-9_])" + re.escape(skill.lower()) + r"(?![a-zA-Z0-9_])"
        if re.search(pattern, text_lower):
            found.append(skill)
            if len(found) >= limit:
                break

    return found
